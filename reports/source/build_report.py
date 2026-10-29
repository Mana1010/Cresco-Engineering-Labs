"""Builds reports/Frappe_UI_Inventory_Report.docx from report_data.py and the components/ screenshots.

Usage (needs python-docx and Pillow, e.g. in a throwaway virtualenv):
    python reports/source/build_report.py components reports/Frappe_UI_Inventory_Report.docx
"""

import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_data import COMPONENTS  # noqa: E402

SHOTS = sys.argv[1]
OUT = sys.argv[2]

BRAND = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x5B, 0x67, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AVAIL_COLORS = {"Yes": ("E3F4E9", "1E7A44"), "No": ("FBE5E6", "B3262E")}

# OOXML requires child elements of tcPr / pPr in a fixed order, so shading must be inserted
# before any of the elements that come after it (appending can make Word reject the file).
TCPR_AFTER_SHD = ("noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark")
PPR_AFTER_SHD = ("tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct",
                 "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
                 "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
                 "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
                 "pPrChange")


def _insert_shd(props, hex_fill, successors):
    for old in props.findall(qn("w:shd")):
        props.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    successor_tags = {qn("w:" + name) for name in successors}
    for child in props:
        if child.tag in successor_tags:
            child.addprevious(shd)
            return
    props.append(shd)


def shade(cell, hex_fill):
    _insert_shd(cell._tc.get_or_add_tcPr(), hex_fill, TCPR_AFTER_SHD)


def shade_paragraph(paragraph, hex_fill):
    _insert_shd(paragraph._p.get_or_add_pPr(), hex_fill, PPR_AFTER_SHD)


def repeat_header(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def set_cell_text(cell, text, *, bold=False, size=8, color=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def mark_available(cell, paragraph, value):
    if value in AVAIL_COLORS:
        fill, ink = AVAIL_COLORS[value]
        shade(cell, fill)
        paragraph.runs[0].bold = True
        paragraph.runs[0].font.color.rgb = RGBColor.from_string(ink)


def fit_image(path, max_w_in, max_h_in):
    """Width that fits the image inside max_w × max_h inches (at 96 dpi), never enlarging it."""
    with Image.open(path) as im:
        w_in, h_in = im.size[0] / 96, im.size[1] / 96
    return Inches(w_in * min(max_w_in / w_in, max_h_in / h_in, 1.0))


def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = BRAND
    return h


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            p.add_run(item[0]).bold = True
            p.add_run(item[1])
        else:
            p.add_run(item)
        p.paragraph_format.space_after = Pt(2)


def add_code(doc, code):
    for line in code.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        shade_paragraph(p, "F4F6F9")
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8)


def small_note(doc, text):
    p = doc.add_paragraph(text)
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = MUTED
    return p


doc = Document()

section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = Cm(29.7), Cm(21.0)
section.left_margin = section.right_margin = Cm(1.5)
section.top_margin = section.bottom_margin = Cm(1.4)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10)

footer_p = section.footer.paragraphs[0]
footer_p.text = "Cresco Engineering Labs Portal — Frappe UI Inventory Report"
footer_p.runs[0].font.size = Pt(8)
footer_p.runs[0].font.color.rgb = MUTED

# ------------------------------------------------------------------ title + intro
title = doc.add_paragraph()
r = title.add_run("Cresco Engineering Labs Portal")
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = BRAND
r = doc.add_paragraph().add_run("Frappe UI Inventory Report — POC components vs Frappe UI")
r.font.size = Pt(14)
r.font.color.rgb = MUTED

meta = doc.add_table(rows=5, cols=2)
meta.style = "Light Grid Accent 1"
for i, (k, v) in enumerate([
    ("Prepared by", "Tristan"),
    ("Date", "2026-10-27"),
    ("Revision", "3 — added component 33 (automation rule cards). Revision 2 added components 25–32 "
                 "(collapsed sidebar, PDF / JSON / code / Excel previews, toast, back link, full Calculation "
                 "Catalogue screen)"),
    ("Evaluated against", "frappe-ui 1.0.2 (latest on npm; components and props checked against the "
                          "frappe/frappe-ui GitHub source and ui.frappe.io docs on 2026-10-27)"),
    ("Input", f"{len(COMPONENTS)} screenshots of POC components (components/ folder)"),
]):
    set_cell_text(meta.cell(i, 0), k, bold=True, size=9)
    set_cell_text(meta.cell(i, 1), v, size=9)
for row in meta.rows:
    row.cells[0].width, row.cells[1].width = Cm(4), Cm(22)

add_heading(doc, "1. Purpose", 1)
doc.add_paragraph(
    "To assess whether Frappe UI (the frappe-ui Vue component library) can deliver the UI/UX of the Cresco "
    "Engineering Labs Portal POC. Each POC component is matched to the closest frappe-ui component, marked as "
    "available or not, and given a recommendation, effort estimate and source code example."
)

add_heading(doc, "2. How to read this report", 1)
add_bullets(doc, [
    ("Available in Frappe UI (Yes / No): ", "Yes = frappe-ui has a dedicated component for this element; it may be "
     "combined with other frappe-ui components and Tailwind layout classes. No = there is no dedicated component, "
     "so a custom component, an external library or a per-instance CSS override is needed."),
    ("Fit: ", "how it is achieved — Out of the box (use as documented); Theme (works, but the POC's brand colours need "
     "a one-time Tailwind theme setting, not per-component code); Compose (several frappe-ui components together); "
     "Custom (new component, external library or CSS override)."),
    ("Effort: ", "Low (under 1 hour), Medium (half a day), High (more than a day)."),
    ("Source code: ", "each component has a Vue 3 example in section 5, written against the frappe-ui 1.0.2 API. "
     "The examples have not been run here: this bench has no frappe-ui setup yet (see finding 1). They show the "
     "intended usage and should be checked once frappe-ui is installed."),
])

yes = sum(1 for c in COMPONENTS if c["available"] == "Yes")
no_names = ", ".join(c.get("short", c["name"]) for c in COMPONENTS if c["available"] == "No")
add_heading(doc, "3. Key findings", 1)
add_bullets(doc, [
    ("Coverage: ", f"{yes} of {len(COMPONENTS)} POC components have a dedicated frappe-ui component. The "
     f"{len(COMPONENTS) - yes} that don't: {no_names}. Within the metrics dashboard, two small parts are also "
     "custom: the coloured quota bars and the two-number review headline."),
    ("1 — Vue version: ", "frappe-ui 1.0.2 requires Vue 3.5 or newer. The Frappe v16 bench in this project bundles "
     "Vue 3.3.9 for its own build, so frappe-ui can't simply be imported into the current portal bundle. The usual "
     "setup is a separate Vite frontend (a frontend/ folder in the app, as Frappe CRM and Helpdesk do), which has "
     "its own Vue version."),
    ("2 — Brand colours: ", "Button themes are gray, blue, green and red; Badge adds amber and violet. The POC's "
     "purple/indigo primary colour and orange action buttons are not built-in themes. They can be set once in the "
     "Tailwind theme, or the client could accept frappe-ui's palette (blue primary)."),
    ("3 — Sidebar: ", "Sidebar supports collapsing to icons natively (covers the collapsible-sidebar task). "
     "A dark sidebar on a light page works by setting data-theme=\"dark\" on the sidebar only; the exact navy and "
     "purple of the POC would need theme overrides."),
    ("4 — No Card / comments / pagination components: ", "frappe-ui has no general Card, no notes/comments "
     "component and no pagination bar. These are small custom components built from frappe-ui primitives "
     "(Badge, Button, TextInput, Select). Pagination matters because the 2026-10-20 meeting requires "
     "first/prev/next/last + rows-per-page on the Tools Registry."),
    ("5 — Progress colour: ", "Progress has no colour option (its bar is always gray), so coloured usage bars need a "
     "small custom bar."),
    ("6 — Charts: ", "the frappe-ui/charts module (NumberCard, BarChart, …) covers the metric cards and bar charts. "
     "It is built on ECharts, which adds to the bundle size."),
    ("7 — Previews: ", "JSON and source-code previews are covered by frappe-ui's code-editor (CodeMirror, read-only, "
     "with JSON and Python highlighting). Two previews are not covered: the PDF viewer (recommended: a Frappe Print "
     "Format rendered to PDF on the server, shown in the browser's own viewer) and the Excel spreadsheet grid "
     "(needs an external library — licences differ — or a simpler inputs-form + results-table design). The "
     "spreadsheet is the single largest item in this report."),
    ("8 — Notifications: ", "toast messages are built in (frappe-ui toast, based on vue-sonner)."),
    ("9 — Look and feel: ", "frappe-ui components follow Frappe's own design language. They match the POC's layout "
     "and behaviour, but colours, spacing and some details (e.g. native select vs frappe-ui popover select) will "
     "look different unless themed. This should be confirmed with the client before committing to frappe-ui."),
])

add_heading(doc, "4. Inventory summary", 1)
small_note(doc, "Rows 2a–2g break down the metrics dashboard screenshot (row 2). "
                "Full screenshots, details and source code for every row are in section 5.")

headers = ["#", "Screenshot", "Component", "Used in", "Frappe UI component", "Available\n(Yes/No)", "Fit",
           "Gap / notes", "Recommendation", "Effort", "Source code"]
widths = [Cm(0.8), Cm(3.2), Cm(2.6), Cm(2.4), Cm(2.8), Cm(1.6), Cm(1.8), Cm(5.4), Cm(3.0), Cm(1.3), Cm(1.6)]
table = doc.add_table(rows=1, cols=len(headers))
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    set_cell_text(cell, h, bold=True, size=8, color=WHITE)
    shade(cell, "1F3A5F")
repeat_header(table.rows[0])


def add_row(values, *, screenshot=None, sub=False):
    row = table.add_row()
    for i, value in enumerate(values):
        cell = row.cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if i == 1:
            cell.text = ""
            if screenshot:
                cell.paragraphs[0].add_run().add_picture(screenshot, width=fit_image(screenshot, 1.15, 1.0))
            else:
                set_cell_text(cell, value, size=7, color=MUTED, italic=True)
            continue
        p = set_cell_text(cell, value, size=7.5 if sub else 8, bold=(i in (0, 2) and not sub))
        if i == 5:
            mark_available(cell, p, value)
    if sub:
        for i in (0, 2, 3, 4, 6, 7, 8, 9, 10):
            shade(row.cells[i], "F7F9FB")


for c in COMPONENTS:
    add_row([c["id"], "", c["name"], c["used_in"], c["frappe"], c["available"], c["fit"], c["gap"],
             c["recommendation"], c["effort"], f"See §5.{c['id']}"], screenshot=os.path.join(SHOTS, c["file"]))
    for s in c.get("subs", []):
        add_row([s["id"], "(part of row 2)", s["name"], "", s["frappe"], s["available"], s["fit"], s["gap"],
                 "", s["effort"], f"See §5.{c['id']}"], sub=True)

for row in table.rows:
    for i, w in enumerate(widths):
        row.cells[i].width = w

# ------------------------------------------------------------------ details
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
add_heading(doc, "5. Component details and source code", 1)
small_note(doc, "Imports: core components from 'frappe-ui', charts from 'frappe-ui/charts', the list from "
                "'frappe-ui/list' and the code viewer from 'frappe-ui/code-editor'. Icons use Lucide, either as strings "
                "(icon=\"lucide-plus\") via frappe-ui's Vite plugin, or as imports from '~icons/lucide/…'.")

for idx, c in enumerate(COMPONENTS):
    if idx:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, f"5.{c['id']}  {c['name']}", 2)
    shot = os.path.join(SHOTS, c["file"])
    doc.add_paragraph().add_run().add_picture(shot, width=fit_image(shot, 9.5, 3.2))
    cap = small_note(doc, f"Screenshot: components/{c['file']}")
    cap.runs[0].font.size = Pt(8)

    info = doc.add_table(rows=0, cols=2)
    info.style = "Light Grid Accent 1"
    for k, v in [
        ("Used in", c["used_in"]),
        ("Frappe UI component", c["frappe"]),
        ("Available in Frappe UI", c["available"]),
        ("Fit / effort", f"{c['fit']} · {c['effort']}"),
        ("Gap / notes", c["gap"]),
        ("Recommendation", c["recommendation"]),
        ("Documentation", "\n".join(c["docs"])),
    ]:
        cells = info.add_row().cells
        set_cell_text(cells[0], k, bold=True, size=9)
        p = set_cell_text(cells[1], v, size=9)
        if k == "Available in Frappe UI":
            mark_available(cells[1], p, v)
    for row in info.rows:
        row.cells[0].width, row.cells[1].width = Cm(4.5), Cm(21.5)

    if c.get("subs"):
        doc.add_paragraph()
        sub_t = doc.add_table(rows=1, cols=5)
        sub_t.style = "Table Grid"
        for i, h in enumerate(["#", "Part", "Frappe UI component", "Available", "Notes"]):
            set_cell_text(sub_t.rows[0].cells[i], h, bold=True, size=8, color=WHITE)
            shade(sub_t.rows[0].cells[i], "1F3A5F")
        for s in c["subs"]:
            cells = sub_t.add_row().cells
            for i, v in enumerate([s["id"], s["name"], s["frappe"], s["available"], s["gap"]]):
                p = set_cell_text(cells[i], v, size=8)
                if i == 3:
                    mark_available(cells[i], p, v)
        for row in sub_t.rows:
            for i, w in enumerate([Cm(1), Cm(6), Cm(4), Cm(2), Cm(13)]):
                row.cells[i].width = w

    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    r = h.add_run("Source code (Vue 3 + frappe-ui 1.0.2)")
    r.bold = True
    r.font.color.rgb = BRAND
    add_code(doc, c["code"])

# ------------------------------------------------------------------ sources
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
add_heading(doc, "6. Sources", 1)
add_bullets(doc, [
    "frappe-ui documentation — https://ui.frappe.io/docs/introduction",
    "frappe-ui source (component list, props, theme tokens) — https://github.com/frappe/frappe-ui",
    "frappe-ui on npm (latest version 1.0.2, peer dependency vue >= 3.5.0) — https://www.npmjs.com/package/frappe-ui",
    "POC screenshots — components/ folder of this project",
])

os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
doc.save(OUT)
print("saved", OUT)
